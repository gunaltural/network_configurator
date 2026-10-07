"""Read SD-WAN planning workbooks locally; no browser CDN dependency."""
import io
import zipfile
import unicodedata


def normalized(value):
    return ' '.join(unicodedata.normalize('NFKC', str(value or '')).strip().replace('İ', 'i').replace('I', 'ı').lower().split())


def parse_sdwan_excel(data):
    if not data or len(data) > 5_000_000:
        raise ValueError('Upload an Excel workbook smaller than 5 MB.')
    try:
        if data.startswith(b'PK'):
            from openpyxl import load_workbook
            with zipfile.ZipFile(io.BytesIO(data)) as archive:
                if len(archive.infolist()) > 1000 or sum(x.file_size for x in archive.infolist()) > 30_000_000:
                    raise ValueError('Excel workbook is too large when expanded.')
                for item in archive.infolist():
                    if item.filename.endswith('.xml'):
                        raw = archive.read(item)
                        if b'<!DOCTYPE' in raw or b'<!ENTITY' in raw:
                            raise ValueError('Unsupported workbook XML.')
            book = load_workbook(io.BytesIO(data), read_only=True, data_only=True, keep_links=False)
            try:
                name = next((n for n in book.sheetnames if normalized(n) == normalized('IP Adres Planlaması')), book.sheetnames[0])
                sheet = book[name]
                if sheet.max_row and sheet.max_row > 5001 or sheet.max_column and sheet.max_column > 150:
                    raise ValueError('Maximum 5,000 planning rows and 150 columns allowed.')
                cells = list(sheet.iter_rows(max_row=5001, max_col=150, values_only=True))
            finally:
                book.close()
        elif data.startswith(b'\xd0\xcf\x11\xe0'):
            import xlrd
            book = xlrd.open_workbook(file_contents=data, on_demand=True)
            try:
                name = next((n for n in book.sheet_names() if normalized(n) == normalized('IP Adres Planlaması')), book.sheet_names()[0])
                sheet = book.sheet_by_name(name)
                if sheet.nrows > 5001 or sheet.ncols > 150:
                    raise ValueError('Maximum 5,000 planning rows and 150 columns allowed.')
                cells = [sheet.row_values(i) for i in range(sheet.nrows)]
            finally:
                book.release_resources()
        else:
            raise ValueError('Select an .xlsx or .xls Excel workbook.')
        # Existing templates may have a title or blank rows before the headers.
        header_index = next((i for i, row in enumerate(cells[:30]) if normalized('Router Hostname') in [normalized(v) for v in row]), None)
        if header_index is None:
            raise ValueError('Router Hostname header was not found in the first 30 rows.')
        def text(value):
            if value is None:
                return ''
            if isinstance(value, float) and value.is_integer():
                return str(int(value))
            return str(value).strip()
        headers = [text(v) for v in cells[header_index]]
        active = [h for h in headers if h]
        if len(set(normalized(h) for h in active)) != len(active):
            raise ValueError('Duplicate column headers are not supported.')
        rows = [{h: text(row[i]) if i < len(row) else '' for i, h in enumerate(headers) if h}
                for row in cells[header_index + 1:] if any(v is not None and str(v).strip() for v in row)]
        if not rows:
            raise ValueError('The planning worksheet has no device rows.')
        return {'sheet': name, 'rows': rows}
    except ValueError:
        raise
    except Exception as exc:
        raise ValueError('Excel workbook could not be read. Check that it is a valid, unencrypted .xlsx or .xls file.') from exc
