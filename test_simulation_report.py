import base64, io, unittest, zipfile
from PIL import Image
from fastapi.testclient import TestClient
from docx import Document
from app import app

class SimulationReportTests(unittest.TestCase):
    def setUp(self):
        self.client=TestClient(app)
        png=io.BytesIO();Image.new('RGB',(400,220),'white').save(png,format='PNG')
        self.scenario=dict(id='test-scenario',createdAt='2026-10-15T13:00:00Z',direction='outbound',clientPrefix='10.10.0.0/24',internetEntry='',preferredGateway='CE-1',activeGateway='CE-2',backupGateway='',failedDevices=['CE-1'],failedLinks=[],baselinePath='Client → Shared LAN → CE-1 → ISP-1 → Internet',currentPath='Client → Shared LAN → CE-2 → ISP-1 → Internet',reachable=True,pathChanged=True,gatewayChanged=True,policies=[dict(connection='wan3 · CE-2 ↔ ISP-1',localPreference=250,prependCount=0)],topologyPng='data:image/png;base64,'+base64.b64encode(png.getvalue()).decode())
    def test_standalone_word_both_languages_and_picture(self):
        for language,title in [('en','Simulation and failover scenarios'),('tr','Simülasyon ve failover senaryoları')]:
            r=self.client.post('/api/network-lab/simulation-word',json={'name':'Pilot failover','language':language,'scenarios':[self.scenario]})
            self.assertEqual(r.status_code,200,r.text[:200] if r.status_code!=200 else '')
            doc=Document(io.BytesIO(r.content));text='\n'.join(p.text for p in doc.paragraphs)+'\n'+'\n'.join(c.text for t in doc.tables for row in t.rows for c in row.cells)
            self.assertIn(title,text);self.assertIn('CE-2',text);self.assertIn('10.10.0.0/24',text)
            self.assertIn('İstemci → Ortak LAN' if language=='tr' else 'Client → Shared LAN',text)
            self.assertEqual(len(doc.inline_shapes),1)
            if language=='tr':self.assertNotIn('Synthetic model results',text)
    def test_empty_or_invalid_scenarios_rejected(self):
        self.assertEqual(self.client.post('/api/network-lab/simulation-word',json={'name':'Pilot','scenarios':[]}).status_code,422)
        bad={**self.scenario,'policies':[{'connection':'wan1','localPreference':100,'prependCount':11}]}
        self.assertEqual(self.client.post('/api/network-lab/simulation-word',json={'name':'Pilot','scenarios':[bad]}).status_code,422)
        bad={**self.scenario,'topologyPng':'data:image/png;base64,ZmFrZQ=='}
        self.assertEqual(self.client.post('/api/network-lab/simulation-word',json={'name':'Pilot','scenarios':[bad]}).status_code,400)
    def test_module_report_preserves_conditional_result_and_omits_bgp_assumptions(self):
        scenario={**self.scenario, 'module':'VPC', 'vendor':'Arista EOS', 'technology':'MLAG',
                  'sourceDevice':'LEAF-1', 'targetDevice':'CORE-1', 'modelStatus':'split-brain',
                  'modelScope':'Physical attachment only', 'findings':['SPLIT-BRAIN / DUAL-ACTIVE: '+('Detailed synchronization impact. '*12)],
                  'roleStates':{'CORE-1':'PRIMARY','CORE-2':'SECONDARY'},
                  'failureSequence':['Peer-keepalive DOWN','Peer-link DOWN'],
                  'failedDevices':['CORE-'+str(i) for i in range(6)]}
        response=self.client.post('/api/network-lab/simulation-word',json={'name':'MLAG scenario','scenarios':[scenario]})
        self.assertEqual(response.status_code,200)
        doc=Document(io.BytesIO(response.content))
        text='\n'.join(p.text for p in doc.paragraphs)+'\n'+'\n'.join(c.text for t in doc.tables for row in t.rows for c in row.cells)
        self.assertIn('split-brain',text)
        self.assertIn('SPLIT-BRAIN / DUAL-ACTIVE',text)
        self.assertEqual(text.count('Detailed synchronization impact.'),12)
        self.assertIn('Arista EOS',text)
        self.assertIn('CORE-5',text)
        self.assertIn('CORE-1: PRIMARY',text)
        self.assertIn('Peer-keepalive DOWN → Peer-link DOWN',text)
        self.assertNotIn('Both CEs',text)
        self.assertNotIn('VRRP',text)

    def test_assets_available(self):
        self.assertEqual(self.client.get('/simulation-report.js').status_code,200)
        self.assertEqual(self.client.get('/simulation-report.js').headers['cache-control'],'no-store')

    def test_word_accepts_all_supported_bgp_policy_counts(self):
        for count in [8,9,16,64]:
            scenario={**self.scenario,'policies':[dict(connection='WAN-'+str(i),localPreference=100+i,prependCount=i%11) for i in range(count)]}
            response=self.client.post('/api/network-lab/simulation-word',json={'name':'Multi-circuit BGP','scenarios':[scenario]})
            self.assertEqual(response.status_code,200)
            doc=Document(io.BytesIO(response.content))
            self.assertEqual(len(doc.tables[-1].rows),count+1)
        scenario['policies'].append(dict(connection='excess',localPreference=100,prependCount=0))
        self.assertEqual(self.client.post('/api/network-lab/simulation-word',json={'scenarios':[scenario]}).status_code,422)

    def test_before_after_images_and_structured_assessment_survive_word_export(self):
        scenario={**self.scenario,'module':'VPC','technology':'vPC','policies':[],
                  'baselineTopologyPng':self.scenario['topologyPng'],
                  'metrics':[dict(metric='Available modeled paths',baseline='2',current='1',unit='paths')],
                  'componentStates':[dict(component='SPINE-2',type='device',baseline='SECONDARY',current='DOWN',detail='Device failure')],
                  'parameterChanges':[dict(parameter='OSPF cost',baseline='10',current='37',reason='Scenario input')],
                  'assessment':{'en':dict(summary='Service path retained with reduced redundancy.',sections=[dict(title='Technical interpretation',paragraphs=['Only one service attachment remains.'])],checks=[dict(criterion='Continuity',expected='Eligible path',status='Not measured')],commands=[dict(command='show vpc brief',purpose='Verify operational roles',execution='Not executed')],references=[])}}
        response=self.client.post('/api/network-lab/simulation-word',json={'name':'Assessment','scenarios':[scenario]})
        self.assertEqual(response.status_code,200,response.text[:200] if response.status_code!=200 else '')
        doc=Document(io.BytesIO(response.content));text='\n'.join(p.text for p in doc.paragraphs)+'\n'+'\n'.join(c.text for t in doc.tables for row in t.rows for c in row.cells)
        self.assertEqual(len(doc.inline_shapes),2)
        for value in ['Baseline topology','Scenario topology','SPINE-2','SECONDARY','DOWN','37','Not executed','show vpc brief','reduced redundancy']:
            self.assertIn(value,text)
        self.assertNotIn('Extra AS copies',text)

    def test_synthetic_bgp_outputs_are_preserved_without_the_generic_command_plan(self):
        scenario={**self.scenario,'syntheticOutputs':{},'assessment':{},'metrics':[dict(metric='Eligible BGP paths',baseline='4',current='1',unit='paths')]}
        for language in ['en','tr']:
            scenario['syntheticOutputs'][language]=[dict(query='model.bgp.neighbors',title='CE-2 BGP',columns=['Device','Peer','State'],rows=[['CE-2 172.16.20.1','ISP-2 172.16.20.2','Available']])]
            scenario['assessment'][language]=dict(summary='CE-2 → ISP-2',sections=[dict(title='Service impact',paragraphs=['CE-1 and ISP-1 down.','No packet-loss percentage, outage duration or convergence time is reported because this scenario contains no measured traffic evidence.'])],commands=[dict(command='show bgp summary',purpose='generic',execution='Not executed')])
            response=self.client.post('/api/network-lab/simulation-word',json={'name':'BGP evidence','language':language,'scenarios':[scenario]})
            self.assertEqual(response.status_code,200,response.text[:200] if response.status_code!=200 else '')
            doc=Document(io.BytesIO(response.content));text='\n'.join(p.text for p in doc.paragraphs)+'\n'+'\n'.join(c.text for t in doc.tables for row in t.rows for c in row.cells)
            for value in ['CE-2 BGP','172.16.20.1','172.16.20.2','4','1']:self.assertIn(value,text)
            self.assertNotIn('No packet-loss percentage',text);self.assertNotIn('show bgp summary',text);self.assertNotIn('Not executed',text)
        scenario['syntheticOutputs']['en'][0]['rows'].append(['mismatched'])
        self.assertEqual(self.client.post('/api/network-lab/simulation-word',json={'scenarios':[scenario]}).status_code,422)

if __name__=='__main__':unittest.main()
