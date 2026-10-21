"""Structured synthetic simulation scenarios shared by standalone and design reports."""
import base64
from io import BytesIO
from typing import Annotated, Literal
from pydantic import BaseModel, Field, model_validator
from docx import Document
from docx.shared import Inches, Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image

Short = Annotated[str, Field(max_length=300)]
class SimulationPolicy(BaseModel):
    connection: Short
    localPreference: int = Field(ge=0, le=4294967295)
    prependCount: int = Field(ge=0, le=10)
class ComponentState(BaseModel):
    component: Short
    type: Short
    baseline: Short
    current: Short
    detail: str = Field(default='', max_length=1500)
class ParameterChange(BaseModel):
    parameter: Short
    baseline: Short
    current: Short
    reason: str = Field(default='', max_length=1500)
class ScenarioMetric(BaseModel):
    metric: Short
    baseline: Short
    current: Short
    unit: Short = ''
class AssessmentSection(BaseModel):
    title: Short
    paragraphs: list[Annotated[str, Field(max_length=5000)]] = Field(default_factory=list, max_length=16)
class AssessmentCheck(BaseModel):
    criterion: Short
    expected: str = Field(max_length=1500)
    status: Short
class AssessmentCommand(BaseModel):
    command: Short
    purpose: str = Field(max_length=1500)
    execution: Short
class AssessmentReference(BaseModel):
    title: Short
    url: str = Field(max_length=1000, pattern=r'^https://')
class ScenarioAssessment(BaseModel):
    summary: str = Field(max_length=4000)
    sections: list[AssessmentSection] = Field(default_factory=list, max_length=12)
    checks: list[AssessmentCheck] = Field(default_factory=list, max_length=16)
    commands: list[AssessmentCommand] = Field(default_factory=list, max_length=40)
    references: list[AssessmentReference] = Field(default_factory=list, max_length=8)
class SyntheticOutput(BaseModel):
    query: Short
    title: Short
    columns: list[Short] = Field(min_length=1, max_length=10)
    rows: list[list[Short]] = Field(default_factory=list, max_length=128)
    @model_validator(mode='after')
    def matching_columns(self):
        if any(len(row)!=len(self.columns) for row in self.rows): raise ValueError('Synthetic output rows must match the column count.')
        return self

class SyntheticCli(BaseModel):
    device: Short
    command: str = Field(max_length=300, pattern=r'^(show|display|get|diagnose)\s')
    deviceAvailable: bool
    output: str = Field(max_length=16000)

class SimulationScenario(BaseModel):
    scenarioName: str = Field(default="", max_length=300)
    scenarioNames: dict[Literal["en", "tr"], Annotated[str, Field(max_length=300)]] = Field(default_factory=dict)
    syntheticCli: list[SyntheticCli] = Field(default_factory=list, max_length=3)
    evidenceVersion: int = Field(default=0, ge=0, le=1)
    pathAvailable: bool | None = None
    forwardingAssurance: Literal['modeled', 'unverified'] | None = None
    syntheticOutputs: dict[Literal['en', 'tr'], Annotated[list[SyntheticOutput], Field(max_length=3)]] = Field(default_factory=dict)
    componentStates: list[ComponentState] = Field(default_factory=list, max_length=1024)
    parameterChanges: list[ParameterChange] = Field(default_factory=list, max_length=256)
    metrics: list[ScenarioMetric] = Field(default_factory=list, max_length=32)
    assessment: dict[Literal['en', 'tr'], ScenarioAssessment] = Field(default_factory=dict)
    baselineTopologyPng: str = Field(default='', max_length=1500000)
    roleStates: dict[Short, Short] = Field(default_factory=dict)
    failureSequence: list[Short] = Field(default_factory=list, max_length=128)
    module: str = Field(default='', max_length=30)
    vendor: Short = ''
    technology: Short = ''
    sourceDevice: Short = ''
    targetDevice: Short = ''
    modelScope: str = Field(default='', max_length=3000)
    modelStatus: Short = ''
    findings: list[Annotated[str, Field(max_length=1500)]] = Field(default_factory=list, max_length=128)
    id: str = Field(max_length=200)
    createdAt: str = Field(max_length=40)
    direction: Literal['outbound', 'inbound']
    clientPrefix: str = Field(max_length=50)
    internetEntry: Short = ''
    preferredGateway: Short
    activeGateway: Short = ''
    backupGateway: Short = ''
    failedDevices: list[Short] = Field(default_factory=list, max_length=128)
    failedLinks: list[Short] = Field(default_factory=list, max_length=512)
    baselinePath: str = Field(max_length=3000)
    currentPath: str = Field(max_length=3000)
    reachable: bool
    pathChanged: bool
    gatewayChanged: bool
    policies: list[SimulationPolicy] = Field(default_factory=list, max_length=64)
    topologyPng: str = Field(default='', max_length=1500000)
class SimulationExport(BaseModel):
    name: str = Field(default='BGP simulation', min_length=1, max_length=120)
    language: Literal['en', 'tr'] = 'en'
    scenarios: list[SimulationScenario] = Field(min_length=1, max_length=40)

def _format_table(table):
    # Keep compact evidence blocks together; allow large inventories to paginate.
    compact = len(table.rows) <= 6
    for index, row in enumerate(table.rows):
        row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
        for cell in row.cells:
            margins=OxmlElement('w:tcMar')
            for side in ['top','bottom','left','right']:
                item=OxmlElement('w:'+side);item.set(qn('w:w'),'55');item.set(qn('w:type'),'dxa');margins.append(item)
            cell._tc.get_or_add_tcPr().append(margins)
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after=Pt(2)
                for run in paragraph.runs: run.font.size=Pt(9)
                paragraph.paragraph_format.keep_with_next=compact and index < len(table.rows)-1

def _table(doc, rows):
    table = doc.add_table(rows=0, cols=2)
    table.style = 'Light Shading Accent 1'
    for label, value in rows:
        cells = table.add_row().cells
        cells[0].text, cells[1].text = str(label), str(value)
    _format_table(table)
    doc.add_paragraph()

def _path(value, tr):
    if not tr: return value
    names = {'Client': 'İstemci', 'Shared LAN': 'Ortak LAN', 'No modeled path': 'Modellenen yol yok', 'No gateway': 'Ağ geçidi yok', 'Source device unavailable': 'Kaynak cihaz kullanılamıyor'}
    return ' → '.join(names.get(x, x) for x in value.split(' → '))

def _matrix(doc, title, headers, rows):
    if not rows: return
    doc.add_heading(title, level=3)
    table = doc.add_table(rows=1, cols=len(headers));table.style='Light Shading Accent 1'
    repeat=OxmlElement('w:tblHeader');table.rows[0]._tr.get_or_add_trPr().append(repeat)
    for cell, value in zip(table.rows[0].cells, headers): cell.text=str(value)
    for values in rows:
        for cell, value in zip(table.add_row().cells, values): cell.text=str(value)
    _format_table(table)

def _metric(value, tr):
    labels={'Selected leaf service attachments':'Seçilen leaf hizmet bağlantıları','Unavailable physical connections':'Kullanılamayan fiziksel bağlantılar','Unavailable logical connections':'Kullanılamayan mantıksal bağlantılar','STP tree roots':'STP root sayısı','STP blocked connections':'STP bloklu bağlantıları','Selected OSPF path cost':'Seçilen OSPF yol maliyeti','Available logical sessions':'Kullanılabilir mantıksal oturumlar','Offered traffic':'Sunulan trafik','Modeled admitted traffic':'Modelde kabul edilen trafik','Configured policy limit':'Yapılandırılmış politika sınırı','Unserved modeled demand':'Modelde karşılanamayan talep','Available WAN transports':'Kullanılabilir WAN taşıyıcıları','Available management connections':'Kullanılabilir yönetim bağlantıları','Leafs with modeled service':'Modelde hizmet yolu bulunan leaf','Interconnect members':'Ara bağlantı üyeleri','Protected service connections':'Korumaya alınan hizmet bağlantıları','Forwarding assurance':'İletim durumu','roots':'root','cost':'maliyet','transports':'taşıyıcı','leafs':'leaf','members':'üye','Selected modeled paths':'Seçilen model yolları','Eligible BGP paths':'Uygun BGP yolları','Usable upstream exits':'Kullanılabilir üst ağ çıkışları','Available eBGP adjacencies':'Kullanılabilir eBGP komşulukları','Available iBGP adjacencies':'Kullanılabilir iBGP komşulukları','Unavailable connections':'Kullanılamayan bağlantılar','Failed devices':'Arızalı cihazlar','Explicitly disabled connections':'Devre dışı bırakılan bağlantılar','paths':'yol','upstreams':'üst ağ','sessions':'komşuluk','connections':'bağlantı','devices':'cihaz','device':'cihaz','connection':'bağlantı','AVAILABLE IN DESIGN':'TASARIMDA KULLANILABİLİR','DOWN':'ARIZALI'}
    return labels.get(value,value) if tr else value

def _assessment(doc, s, tr):
    a=s.get('assessment',{}).get('tr' if tr else 'en')
    if not a: return
    doc.add_heading('Senaryo değerlendirmesi' if tr else 'Scenario assessment',level=3)
    doc.add_paragraph(a['summary'])
    bgp=not s.get('module') or s.get('module')=='BGP'
    outputs=s.get('syntheticOutputs',{}).get('tr' if tr else 'en',[])
    cli=s.get('syntheticCli',[])
    rich=bool(outputs or cli)
    if cli: outputs=[x for x in outputs if x['query'] in ['model.bgp.routes','model.pair.leaves','model.endpoint.paths']]
    _matrix(doc,'Karşılaştırmalı göstergeler' if tr else 'Comparative metrics',
            ['Gösterge','Başlangıç','Senaryo','Birim'] if tr else ['Metric','Baseline','Scenario','Unit'],
            [[_metric('Selected modeled paths' if bgp and not rich and k=='metric' and x[k]=='Available modeled paths' else x[k],tr) for k in ['metric','baseline','current','unit']] for x in s.get('metrics',[])])
    _matrix(doc,'Bileşenlerin önce ve sonra durumu' if tr else 'Component state before and after',
            ['Bileşen','Tür','Başlangıç','Senaryo','Ayrıntı'] if tr else ['Component','Type','Baseline','Scenario','Detail'],
            [[_metric(x[k],tr) for k in ['component','type','baseline','current','detail']] for x in s.get('componentStates',[]) if not rich or any(word in x['current'] for word in ['DOWN','BLOCKED','UNVERIFIED','RECOVERY']) or x['baseline']!=x['current'] and any(word in x['baseline']+' '+x['current'] for word in ['PRIMARY','SECONDARY','ACTIVE','STANDBY','MASTER','BACKUP','ROOT'])])
    _matrix(doc,'Tasarım ve senaryo parametreleri' if tr else 'Design and scenario parameters',
            ['Parametre','Tasarım','Senaryo','Gerekçe'] if tr else ['Parameter','Design','Scenario','Reason'],
            [[x[k] for k in ['parameter','baseline','current','reason']] for x in s.get('parameterChanges',[]) if not rich or x['baseline']!=x['current']])
    if cli:
        doc.add_heading('Sentetik show çıktıları' if tr else 'Synthetic show outputs',level=3)
        doc.add_paragraph('Vendor komut biçimiyle sunulan model alanları; anlık cihaz CLI kaydı değildir.' if tr else 'Model fields presented with vendor command names; not a captured device CLI transcript.')
        for item in cli:
            doc.add_heading(item['device']+' · '+item['command'],level=3)
            prompt='<'+item['device']+'>' if item['command'].startswith('display ') else item['device']+'#'
            paragraph=doc.add_paragraph()
            paragraph.paragraph_format.keep_together=True
            run=paragraph.add_run(prompt+' '+item['command']+'\n'+item['output']);run.font.name='Consolas';run.font.size=Pt(8)
    if outputs:
        doc.add_heading(('Sentetik BGP çıktıları' if tr else 'Synthetic BGP outputs') if bgp else ('Modüle özel sonuçlar' if tr else 'Module-specific results'),level=3)
        for output in outputs: _matrix(doc,output['title'],output['columns'],output['rows'])
    omitted={'Verification procedure','Scope and evidence limitations','Doğrulama prosedürü','Kapsam ve kanıt sınırları','Scenario objective','Baseline and design context','Applied events and parameters','Observed model behavior','Senaryonun amacı','Başlangıç ve tasarım bağlamı','Uygulanan olaylar ve parametreler','Gözlenen model davranışı'}
    for section in a.get('sections',[]):
        if section['title'] in omitted: continue
        doc.add_heading(section['title'],level=3)
        for paragraph in section['paragraphs']:
            if not paragraph.startswith(('No packet-loss percentage','Bu senaryoda ölçülmüş trafik kanıtı')):doc.add_paragraph(paragraph)
    _matrix(doc,'Kontrol ölçütleri ve kanıt durumu' if tr else 'Verification criteria and evidence status',
            ['Ölçüt','Beklenen','Kanıt durumu'] if tr else ['Criterion','Expected','Evidence status'],
            [[x[k] for k in ['criterion','expected','status']] for x in ([] if bgp or rich else a.get('checks',[]))])
    _matrix(doc,'Doğrulama komut planı' if tr else 'Verification command plan',
            ['Komut','Amaç','Yürütme'] if tr else ['Command','Purpose','Execution'],
            [[x[k] for k in ['command','purpose','execution']] for x in ([] if bgp or rich else a.get('commands',[]))])
    for ref in a.get('references',[]): doc.add_paragraph(ref['title']+' — '+ref['url'])

def append_simulation_scenarios(doc, scenarios, language='en'):
    if not scenarios: return
    tr = language == 'tr'
    doc.add_heading('Simülasyon ve failover senaryoları' if tr else 'Simulation and failover scenarios', level=1)
    doc.add_paragraph('Sentetik senaryo analizi · Tasarım ve arıza modelinden üretilmiştir.' if tr else 'Synthetic model results · Generated from the design and failure model.')
    if any(not (s.model_dump() if isinstance(s, BaseModel) else s).get("module") for s in scenarios):
        doc.add_paragraph('Her iki CE ortak istemci LAN’ına doğrudan bağlıdır. VRRP istemci çıkış ağ geçidini seçer; gelen trafik BGP politikasına göre yedek CE’den de geçebilir. Aktif CE arızasında sağlıklı yedek devralır; tercih edilen CE geri geldiğinde tekrar aktif olur. Internet bulutu modellenen üst ağ sınırını gösterir.' if tr else 'Both CEs directly attach to the shared client LAN. VRRP selects the client outbound gateway; inbound BGP traffic may enter through the backup CE. Device failure promotes the healthy backup and restoration preempts to the preferred CE. The Internet cloud marks the modeled upstream boundary.')
    for index, raw in enumerate(scenarios, 1):
        s = raw.model_dump() if isinstance(raw, BaseModel) else raw
        doc.add_heading(('Senaryo ' if tr else 'Scenario ') + str(index) + (' · ' + (s.get('scenarioNames', {}).get('tr' if tr else 'en') or s.get('scenarioName')) if s.get('scenarioNames', {}).get('tr' if tr else 'en') or s.get('scenarioName') else ''), level=2)
        if s.get('module'):
            _table(doc, [
                ('Modül / platform' if tr else 'Module / platform', s.get('technology', s['module']) + ' · ' + s.get('vendor', '')),
                ('Kaynak / hedef' if tr else 'Source / destination', s.get('sourceDevice','') + ' → ' + s.get('targetDevice','')),
                ('Model sonucu' if tr else 'Model result', s.get('modelStatus','')),
                ('Model kapsamı' if tr else 'Model scope', s.get('modelScope','')),
                ('Bulgular' if tr else 'Findings', '; '.join(s.get('findings',[])) or '—'),
                ('Operasyonel roller' if tr else 'Operational roles', '; '.join(k+': '+v for k,v in s.get('roleStates',{}).items()) or '—'),
                ('Arıza sırası' if tr else 'Failure sequence', ' → '.join(s.get('failureSequence',[])) or '—'),
            ])
        outcome = ('Erişilebilir' if s['reachable'] else 'Yol yok') if tr else ('Reachable' if s['reachable'] else 'No path')
        if s.get('module'):
            _table(doc, [
                ('Kayıt zamanı' if tr else 'Recorded at', s['createdAt']),
                ('Arızalı cihazlar' if tr else 'Failed devices', ', '.join(s['failedDevices']) or '—'),
                ('Arızalı bağlantılar' if tr else 'Failed connections', ', '.join(s['failedLinks']) or '—'),
                ('Başlangıç yolu' if tr else 'Baseline path', _path(s['baselinePath'],tr)),
                ('Senaryo yolu' if tr else 'Scenario path', _path(s['currentPath'],tr)),
            ])
        else:
            _table(doc, [
                ('Kayıt zamanı' if tr else 'Recorded at', s['createdAt']),
                ('Trafik yönü' if tr else 'Traffic direction', ('Çıkış · İstemci → Internet' if s['direction']=='outbound' else 'Giriş · Internet → İstemci') if tr else ('Outbound · Client → Internet' if s['direction']=='outbound' else 'Inbound · Internet → Client')),
                ('İstemci LAN prefix’i' if tr else 'Client LAN prefix', s['clientPrefix']),
                ('Internet girişi' if tr else 'Internet entry', s.get('internetEntry') or '—'),
                ('Tercih edilen / aktif / yedek VRRP' if tr else 'Preferred / active / backup VRRP', ' / '.join(s.get(k) or '—' for k in ['preferredGateway','activeGateway','backupGateway'])),
                ('Arızalı cihazlar' if tr else 'Failed devices', ', '.join(s['failedDevices']) or '—'),
                ('Devre dışı bağlantılar' if tr else 'Disabled connections', ', '.join(s['failedLinks']) or '—'),
                ('Arızasız başlangıç yolu' if tr else 'No-fault baseline path', _path(s['baselinePath'],tr)),
                ('Senaryo yolu' if tr else 'Scenario path', _path(s['currentPath'],tr)),
                ('Model sonucu' if tr else 'Model result', outcome),
                ('Yol değişimi / ağ geçidi değişimi' if tr else 'Path change / gateway change', ' / '.join(('Evet' if flag else 'Hayır') if tr else ('Yes' if flag else 'No') for flag in [s['pathChanged'],s['gatewayChanged']])),
            ])
        for field, caption in [('baselineTopologyPng','Başlangıç topolojisi' if tr else 'Baseline topology'),('topologyPng','Senaryo topolojisi' if tr else 'Scenario topology')]:
            if not s.get(field): continue
            encoded=s[field]
            if not encoded.startswith('data:image/png;base64,'): raise ValueError('Invalid simulation topology image.')
            data=base64.b64decode(encoded.split(',',1)[1],validate=True)
            if not data.startswith(b'\x89PNG\r\n\x1a\n'): raise ValueError('Invalid simulation topology image.')
            image=Image.open(BytesIO(data))
            if image.width>2400 or image.height>2400: raise ValueError('Simulation topology image is too large.')
            doc.add_heading(caption,level=3)
            doc.add_picture(BytesIO(data),width=Inches(6.5))
        _assessment(doc,s,tr)
        if s['policies'] and not s.get('syntheticOutputs'):
            doc.add_heading('Uygulanan sentetik politikalar' if tr else 'Applied synthetic policies', level=3)
            table=doc.add_table(rows=1,cols=3);table.style='Light Shading Accent 1'
            for cell, label in zip(table.rows[0].cells, ['Bağlantı' if tr else 'Connection','Local Preference','Ek AS kopyası' if tr else 'Extra AS copies']):cell.text=label
            for p in s['policies']:
                cells=table.add_row().cells
                for cell,value in zip(cells,[p['connection'],p['localPreference'],p['prependCount']]):cell.text=str(value)

def build_simulation_docx(payload):
    doc=Document();doc.sections[0].left_margin=doc.sections[0].right_margin=Inches(.7);doc.styles['Normal'].font.name='Arial';doc.styles['Normal'].font.size=Pt(10)
    doc.add_heading(payload['name'],0)
    append_simulation_scenarios(doc,payload['scenarios'],payload.get('language','en'))
    out=BytesIO();doc.save(out);out.seek(0);return out
