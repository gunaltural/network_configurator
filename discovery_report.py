"""Standalone brownfield reports; no greenfield inventory or report mutation."""
from io import BytesIO
import math
from pathlib import Path
from datetime import datetime, timezone
from xml.sax.saxutils import escape
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from reportlab.lib import colors
from reportlab.lib.pagesizes import landscape, letter
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, Image as PDFImage, PageBreak
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
import reportlab

LABELS={
'en':dict(title='Network discovery report',intro='Observed devices and port connections in the existing network. This brownfield document is independent of the planned design and Inventory report.',summary='Discovery scope and collection summary',topology='Observed network topology',devices='Discovered devices',links='Connection details',sources='Collection sources',warnings='Collection and parsing notes',project='Project',engineer='Engineer',scope='Scope',generated='Report generated',coverage='Coverage',method='Collection methods',missing='Not collected',one='One-sided observation',both='Both ends observed',limitation='Neighbor advertisements establish observed adjacencies. They do not verify traffic health, redundancy or application reachability. An absent neighbor is not proof of a failed connection.',maplimit='The diagram shows up to 200 devices. The tables include every discovered device and connection.',neighbor='Neighbor advertisement only',direct='Collected or imported locally',provenance='Interface evidence',nointerfaces='No additional interface output was collected for this endpoint.'),
 'tr':dict(title='Ağ keşif raporu',intro='Mevcut ağda gözlemlenen cihazlar ve port bağlantıları. Bu brownfield dokümanı planlanan tasarımdan ve Inventory raporundan bağımsızdır.',summary='Keşif kapsamı ve toplama özeti',topology='Gözlemlenen ağ topolojisi',devices='Keşfedilen cihazlar',links='Bağlantı detayları',sources='Veri toplama kaynakları',warnings='Toplama ve ayrıştırma notları',project='Proje',engineer='Mühendis',scope='Kapsam',generated='Rapor oluşturma zamanı',coverage='Kapsama',method='Veri toplama yöntemleri',missing='Toplanmadı',one='Tek uçtan gözlemlendi',both='İki uçtan gözlemlendi',limitation='Komşuluk bildirimleri gözlemlenen bağlantıları gösterir. Trafik sağlığını, yedekliliği veya uygulama erişimini doğrulamaz. Bir komşunun görülmemesi bağlantı arızasının kanıtı değildir.',maplimit='Şema en fazla 200 cihazı gösterir. Tablolar keşfedilen bütün cihazları ve bağlantıları içerir.',neighbor='Yalnız komşuluk bildirimi',direct='Cihazdan toplandı veya yerel çıktı aktarıldı',provenance='Arayüz verisinin kaynağı',nointerfaces='Bu uç için ek arayüz çıktısı toplanmadı.')}


def safe(value):
    return ''.join(c for c in str(value or '') if c in '\n\t' or ord(c)>=32)[:4000]


def topology_png(graph, positions):
    """Bounded render using observed nodes/edges and the current map positions."""
    nodes=graph['devices'][:200];ids={d['id'] for d in nodes};points={}
    cols=max(2,min(6,math.ceil(math.sqrt(max(1,len(nodes))))))
    for i,d in enumerate(nodes):
        p=positions.get(d['id'],{})
        x,y=p.get('x'),p.get('y')
        if not isinstance(x,(int,float)) or not isinstance(y,(int,float)) or not math.isfinite(x) or not math.isfinite(y) or abs(x)>20000 or abs(y)>20000:
            x,y=110+(i%cols)*220,65+(i//cols)*110
        points[d['id']]=(x,y)
    if not nodes:return None
    left=min(p[0] for p in points.values())-100;top=min(p[1] for p in points.values())-60
    width=max(p[0] for p in points.values())-left+100;height=max(p[1] for p in points.values())-top+80
    parallel={}
    for edge in graph['links']:
        pair=tuple(sorted((edge['a'],edge['b'])))
        parallel[pair]=parallel.get(pair,0)+1
    height += max(0,max(parallel.values(),default=1)-1)*30
    scale=min(2,2200/max(1,width),2200/max(1,height));im=Image.new('RGB',(max(1,int(width*scale)),max(1,int(height*scale))),'#f7fbfd');draw=ImageDraw.Draw(im)
    fontpath=Path(reportlab.__file__).parent/'fonts'/'Vera.ttf'
    font=ImageFont.truetype(str(fontpath),max(10,int(10*scale)))
    def xy(p):return ((p[0]-left)*scale,(p[1]-top)*scale)
    groups={}
    for e in graph['links']:
        if e['a'] in ids and e['b'] in ids:groups.setdefault(tuple(sorted((e['a'],e['b']))),[]).append(e)
    for group in groups.values():
        for index,e in enumerate(group):
            a,b=xy(points[e['a']]),xy(points[e['b']]);mid=((a[0]+b[0])/2,(a[1]+b[1])/2+index*30*scale)
            for start,end in [(a,mid),(mid,b)]:
                if e['state']=='Both ends observed':draw.line([start,end],fill='#4889a5',width=max(1,int(2*scale)))
                else:
                    steps=max(1,math.ceil(math.dist(start,end)/max(1,6*scale)))
                    for step in range(0,steps,2):
                        u=step/steps;v=min(1,(step+1)/steps)
                        draw.line([(start[0]+(end[0]-start[0])*u,start[1]+(end[1]-start[1])*u),(start[0]+(end[0]-start[0])*v,start[1]+(end[1]-start[1])*v)],fill='#4889a5',width=max(1,int(2*scale)))
            # Same endpoints are retained for parallel links, each label has its own offset.
            draw.text((mid[0],mid[1]-9*scale),safe(e['aPort']),fill='#264e69',font=font,anchor='mm')
            draw.text((mid[0],mid[1]+9*scale),safe(e['bPort']),fill='#264e69',font=font,anchor='mm')
    for d in nodes:
        x,y=xy(points[d['id']]);rect=(x-77*scale,y-39*scale,x+77*scale,y+39*scale)
        fill='#a34e58' if d['vendor'].startswith('Huawei') else '#24736a' if d['vendor']=='Arista EOS' else '#1b628b'
        draw.rounded_rectangle(rect,radius=7*scale,fill=fill)
        text=safe(d['hostname']);lines=[text[i:i+24] for i in range(0,len(text),24)][:3]
        for i,line in enumerate(lines):draw.text((x,y-24*scale+i*11*scale),line,fill='white',font=font,anchor='mm')
        ip=d['managementAddresses'][0] if d['managementAddresses'] else '-'
        draw.text((x,y+24*scale),safe(ip),fill='white',font=font,anchor='mm')
    out=BytesIO();im.save(out,format='PNG');out.seek(0);return out


def report_sections(graph,sources,data):
    tr=data.get('language')=='tr';l=LABELS['tr' if tr else 'en'];miss=l['missing'];byid={d['id']:d for d in graph['devices']}
    def origin(value):return {'CLI import':'CLI çıktısı aktarımı','SSH inventory / neighbors':'SSH envanter ve komşuluk toplama','SSH interface collection':'SSH arayüz toplama'}.get(value,value) if tr else value
    def state(e):return l['both'] if e['state']=='Both ends observed' else l['one']
    meta=[[l['project'],data['name']],[l['generated'],datetime.now(timezone.utc).isoformat(timespec='seconds')],[l['coverage'],f"{len(graph['devices'])} {'cihaz' if tr else 'devices'} / {len(graph['links'])} {'bağlantı' if tr else 'connections'}"],[l['method'],' / '.join(sorted({origin(s['origin']) for s in sources}))]]
    if data.get('engineer'):meta.append([l['engineer'],data['engineer']])
    deviceheaders=['Cihaz','Yönetim IP','Platform','Model','Seri no','Yazılım'] if tr else ['Device','Management IP','Platform','Model','Serial','Software']
    devicerows=[[d['hostname'],', '.join(d['managementAddresses']) or miss,d['vendor'] or miss,d['model'] or miss,d['serial'] or miss,d['softwareVersion'] or miss] for d in graph['devices']]
    links=[]
    for e in graph['links']:
        endpointrows=[];notes=[]
        for side in ('a','b'):
            d=byid[e[side]];r=e.get(side+'Details') or {}
            endpointrows.append([d['hostname'],e[side+'Port'],' / '.join(filter(None,[r.get('status'),r.get('protocol')])) or miss,r.get('speed') or miss,' / '.join(filter(None,[r.get('mode'),r.get('vlan')])) or miss,r.get('portChannel') or miss])
            if r.get('description'):notes.append(d['hostname']+': '+r['description'])
            if r.get('allowedVlans'):notes.append(d['hostname']+(' izinli VLAN: ' if tr else ' allowed VLANs: ')+r['allowedVlans'])
            for o in r.get('evidence',[]):notes.append(d['hostname']+' / '+o['command']+' / '+o.get('observedAt',''))
        title=byid[e['a']]['hostname']+' / '+e['aPort']+' - '+byid[e['b']]['hostname']+' / '+e['bPort']
        details=state(e)+' / '+' / '.join(e['protocols'])
        headers=['Cihaz','Port','Durum','Hız','Mod / VLAN','Port Channel'] if tr else ['Device','Port','State','Speed','Mode / VLAN','Port Channel']
        links.append((title,details,headers,endpointrows,notes))
    sourceheaders=['Cihaz','Protokol','Yöntem','Gözlem zamanı'] if tr else ['Device','Protocol','Method','Observed at']
    sourcerows=[[s['deviceName'],s['protocol'].upper(),origin(s['origin']),s['observedAt']] for s in sources]
    return l,meta,deviceheaders,devicerows,links,sourceheaders,sourcerows


def docx_table(doc,headers,rows,widths):
    table=doc.add_table(rows=1,cols=len(headers));table.autofit=False
    for cell,text,width in zip(table.rows[0].cells,headers,widths):cell.text=safe(text);cell.width=Inches(width)
    repeat=OxmlElement('w:tblHeader');table.rows[0]._tr.get_or_add_trPr().append(repeat)
    for row in rows:
        cells=table.add_row().cells
        for cell,text,width in zip(cells,row,widths):cell.text=safe(text);cell.width=Inches(width)
    for column,width in zip(table.columns,widths):column.width=Inches(width)
    for i,row in enumerate(table.rows):
        for cell in row.cells:
            pr=cell._tc.get_or_add_tcPr();shd=OxmlElement('w:shd');shd.set(qn('w:fill'),'DCE8F0' if i==0 else 'F4F7F9' if i%2==0 else 'FFFFFF');pr.append(shd)
            borders=OxmlElement('w:tcBorders')
            for edge in ('top','left','bottom','right'):
                element=OxmlElement('w:'+edge);element.set(qn('w:val'),'single');element.set(qn('w:sz'),'4');element.set(qn('w:color'),'D9D9D9');borders.append(element)
            pr.append(borders)
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after=Pt(4);paragraph.paragraph_format.space_before=Pt(4)
                for run in paragraph.runs:run.font.size=Pt(9);run.bold=i==0
    gap=doc.add_paragraph();gap.paragraph_format.line_spacing=Pt(1);gap.paragraph_format.space_after=Pt(5)


def build_discovery_report(graph,sources,data,format):
    l,meta,dh,dr,links,sh,sr=report_sections(graph,sources,data);image=topology_png(graph,data.get('positions',{}));out=BytesIO()
    if format=='word':
        doc=Document();sec=doc.sections[0];sec.page_width=Inches(11);sec.page_height=Inches(8.5);sec.top_margin=sec.bottom_margin=Inches(.6);sec.left_margin=sec.right_margin=Inches(.65)
        for name in ('Normal','Title','Heading 1','Heading 2'):
            style=doc.styles[name];style.font.name='Calibri';style.font.color.rgb=RGBColor(0,0,0)
        doc.styles['Normal'].font.size=Pt(10)
        doc.styles['Normal'].paragraph_format.space_after=Pt(5)
        doc.styles['Title'].font.size=Pt(24)
        for name in ('Title','Heading 1','Heading 2'):
            doc.styles[name].paragraph_format.space_before=Pt(8);doc.styles[name].paragraph_format.space_after=Pt(5)
        for style in doc.styles:
            for border in list(style.element.iter(qn('w:pBdr'))):border.getparent().remove(border)
        doc.add_paragraph(l['title'],'Title');doc.add_paragraph(l['intro']);doc.add_heading(l['summary'],1);docx_table(doc,['Alan' if data.get('language')=='tr' else 'Field','Değer' if data.get('language')=='tr' else 'Value'],meta,[2,7.7])
        if data.get('scope'):doc.add_heading(l['scope'],1);doc.add_paragraph(safe(data['scope']))
        doc.add_paragraph(l['limitation']);doc.add_heading(l['topology'],1)
        if image:
            w,h=Image.open(image).size;image.seek(0);doc.add_picture(image,width=Inches(min(9.5,1.8*w/h)))
        if len(graph['devices'])>200:doc.add_paragraph(l['maplimit'])
        doc.add_heading(l['devices'],1);docx_table(doc,dh,dr,[2,1.9,1.5,1.5,1.3,1.5])
        doc.add_heading(l['links'],1)
        for title,details,headers,rows,notes in links:
            doc.add_heading(safe(title),2);detail=doc.add_paragraph(details);detail.paragraph_format.keep_with_next=True;docx_table(doc,headers,rows,[2,1.5,1.4,1,2.3,1.5])
            for note in notes:doc.add_paragraph(safe(note))
        doc.add_heading(l['sources'],1);docx_table(doc,sh,sr,[2.4,1,2.6,3.7])
        if graph['warnings']:
            doc.add_heading(l['warnings'],1)
            for warning in graph['warnings']:doc.add_paragraph(safe(warning))
        footer=sec.footer.paragraphs[0];footer.add_run(l['title']+' | ')
        field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
        doc.save(out)
    else:
        fontdir=Path(reportlab.__file__).parent/'fonts'
        for name,file in [('Discovery','Vera.ttf'),('DiscoveryBold','VeraBd.ttf')]:
            if name not in pdfmetrics.getRegisteredFontNames():pdfmetrics.registerFont(TTFont(name,str(fontdir/file)))
        styles=getSampleStyleSheet()
        for name in ('Normal','Title','Heading1','Heading2'):
            styles[name].fontName='DiscoveryBold' if name!='Normal' else 'Discovery';styles[name].textColor=colors.black
        styles['Normal'].fontSize=9;styles['Normal'].leading=13;styles['Heading2'].fontSize=11;styles['Title'].alignment=0;styles['Title'].fontSize=24;styles['Title'].leading=30;styles['Title'].spaceAfter=10;styles['Heading1'].keepWithNext=True;styles['Heading2'].keepWithNext=True
        body=[]
        def para(text,style='Normal'):return Paragraph(escape(safe(text)).replace('\n','<br/>'),styles[style])
        def table(headers,rows,widths):
            values=[[para(c) for c in headers]]+[[para(c) for c in row] for row in rows]
            t=Table(values,colWidths=widths,repeatRows=1,hAlign='LEFT');t.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),colors.HexColor('#dce8f0')),('ROWBACKGROUNDS',(0,1),(-1,-1),[colors.white,colors.HexColor('#f4f7f9')]),('GRID',(0,0),(-1,-1),.4,colors.HexColor('#d9d9d9')),('VALIGN',(0,0),(-1,-1),'MIDDLE'),('LEFTPADDING',(0,0),(-1,-1),7),('RIGHTPADDING',(0,0),(-1,-1),7),('TOPPADDING',(0,0),(-1,-1),6),('BOTTOMPADDING',(0,0),(-1,-1),6)]));body.extend([t,Spacer(1,12)])
        body.extend([para(l['title'],'Title'),para(l['intro']),Spacer(1,12)]);body.append(para(l['summary'],'Heading1'));table(['Alan' if data.get('language')=='tr' else 'Field','Değer' if data.get('language')=='tr' else 'Value'],meta,[144,554])
        if data.get('scope'):body.extend([para(l['scope'],'Heading1'),para(data['scope'])])
        body.extend([para(l['limitation']),para(l['topology'],'Heading1')])
        if image:
            w,h=Image.open(image).size;image.seek(0);scale=min(698/w,160/h);body.append(PDFImage(image,width=w*scale,height=h*scale))
        if len(graph['devices'])>200:body.append(para(l['maplimit']))
        body.append(para(l['devices'],'Heading1'));table(dh,dr,[144,137,108,108,94,107])
        body.append(para(l['links'],'Heading1'))
        for title,details,headers,rows,notes in links:
            body.extend([para(title,'Heading2'),para(details),Spacer(1,6)]);table(headers,rows,[144,108,101,72,166,107])
            for note in notes:body.append(para(note))
        body.append(para(l['sources'],'Heading1'));table(sh,sr,[173,72,187,266])
        if graph['warnings']:
            body.append(para(l['warnings'],'Heading1'))
            body.extend(para(w) for w in graph['warnings'])
        def footer(canvas,doc):
            canvas.setFont('Discovery',8);canvas.drawString(47,22,l['title']);canvas.drawRightString(745,22,str(doc.page))
        SimpleDocTemplate(out,pagesize=landscape(letter),leftMargin=47,rightMargin=47,topMargin=40,bottomMargin=40,title=l['title']).build(body,onFirstPage=footer,onLaterPages=footer)
    out.seek(0);return out
