"""Regression checks for reports and read-only SSH; no real devices contacted."""
import base64
import io
import unittest
from unittest.mock import patch
from PIL import Image
from docx import Document
from fastapi.testclient import TestClient
import app


def device(id, tier='upper', index=1, **extra):
    return dict(id=id,tier=tier,index=index,hostname=id,model='',serial='',modelSource='',serialSource='',**extra)


def project():
    return dict(name='Combined audit',language='tr',vendor='Cisco NX-OS',architecture='module',technology='Combined · vPC + System and Management + BGP',techPlacement='upper',upperCount=2,lowerCount=0,scope='',devices=[device('vpc-upper-1'),device('basic-upper-1',index=2)],links=[dict(a='vpc-upper-1',b='basic-upper-1',enabled=True,upperPort='Eth1/1',lowerPort='Eth1/2',speed='100G')],inventoryDeviceIds=[])


class ApplicationAuditTests(unittest.TestCase):
    def setUp(self):
        self.client=TestClient(app.app)

    def test_combined_word_and_inventory_transfer(self):
        p=project(); p['devices'].append(device('inventory-12345678-1234-1234-1234-123456789012',inventoryOnly=True,observedHostname='ACTUAL-SW',softwareVersion='10.5(3)'))
        p['inventoryDeviceIds']=[p['devices'][-1]['id']]
        for lang in ['tr','en']:
            p['language']=lang;r=self.client.post('/api/reporting/word',json=p)
            self.assertEqual(r.status_code,200,r.text if r.status_code!=200 else '')
            self.assertIn('ACTUAL-SW',' '.join(c.text for t in Document(io.BytesIO(r.content)).tables for row in t.rows for c in row.cells))

    def test_empty_inventory_is_not_claimed_complete(self):
        p=project()
        for language in ['tr','en']:
            p['language']=language
            r=self.client.post('/api/reporting/word',json=p)
            doc=Document(io.BytesIO(r.content));text=' '.join(x.text for x in doc.paragraphs)
            self.assertIn('Henüz envanter kaydı yok' if language=='tr' else 'No inventory records yet',text)
            self.assertNotIn('Tüm cihaz modelleri' if language=='tr' else 'All device models',text)

    def test_combined_word_all_module_topologies(self):
        p=project();p['moduleReports']=[]
        for title,color in [('vPC','blue'),('System and Management','green')]:
            stream=io.BytesIO();Image.new('RGB',(100,50),color).save(stream,format='PNG')
            p['moduleReports'].append(dict(title=title,vendor='Cisco NX-OS',topologyPng='data:image/png;base64,'+base64.b64encode(stream.getvalue()).decode()))
        r=self.client.post('/api/reporting/word',json=p)
        self.assertEqual(r.status_code,200)
        self.assertEqual(len(Document(io.BytesIO(r.content)).inline_shapes),2)

    def test_combined_topology_integrity(self):
        for mutate in [lambda p:p['devices'].append(p['devices'][0]),lambda p:p['links'][0].update(b='missing'),lambda p:p.update(upperCount=3),lambda p:p.update(inventoryDeviceIds=['missing'])]:
            p=project();mutate(p)
            self.assertEqual(self.client.post('/api/reporting/word',json=p).status_code,400)

    def test_large_combined_topology(self):
        p=project();p['devices']=[device(f'module-device-{i}',index=i) for i in range(1,31)];p['upperCount']=30;p['links']=[]
        self.assertEqual(self.client.post('/api/reporting/word',json=p).status_code,200)

    def test_every_catalog_command_is_accepted(self):
        for platform,groups in app.LIVE_COMMANDS.items():
            for commands in groups.values():
                for command in commands:self.assertTrue(app.live_command_allowed(platform,command),(platform,command))
        for platform in app.LIVE_COMMANDS:
            for command in ['configure terminal','show version; reload','show version | redirect file','show version && reload']:
                self.assertFalse(app.live_command_allowed(platform,command))
        self.assertEqual(len(app.live_command_list('Cisco NX-OS','\n'.join(['show version']*101))),101)

    def test_fortigate_connection_test_adapter(self):
        with patch.object(app,'MOCK_SSH',True),patch.object(app,'resolve_ssh_target',return_value='8.8.8.8'):
            r=self.client.post('/api/device/test',json=dict(target='test.example',platform='FortiGate',username='test'))
            self.assertEqual(r.status_code,200,r.text)
            self.assertIn('fortinet',str(r.json()))

    def test_live_command_failure_isolation(self):
        class Connection:
            def send_command(self,command,**kwargs):
                if command=='show clock':raise RuntimeError('Simulated failure')
                return 'TEST OUTPUT'
            def disconnect(self):self.disconnected=True
        conn=Connection()
        with patch.object(app,'MOCK_SSH',False),patch.object(app,'resolve_ssh_target',return_value='8.8.8.8'),patch.object(app,'connect',return_value=(conn,'cisco_nxos')):
            r=self.client.post('/api/device/show',json=dict(target='test.example',platform='Cisco NX-OS',username='test',password='synthetic',command='show version\nshow clock\nshow inventory'))
        self.assertEqual(r.status_code,200)
        self.assertEqual([x['ok'] for x in r.json()['results']],[True,False,True]);self.assertTrue(conn.disconnected)
        self.assertEqual([x['output'] for x in r.json()['results']],['TEST OUTPUT','Simulated failure','TEST OUTPUT'])
        self.assertTrue(all(not x['truncated'] for x in r.json()['results']))

    def test_malformed_excel_shared_string(self):
        from inventory_excel import parse_inventory_xlsx
        from test_inventory_excel import fixture
        import zipfile
        data=fixture([['Hostname','Serial Number','Software Version','Product Model'],['SW','SN','1','Model']]);out=io.BytesIO()
        with zipfile.ZipFile(io.BytesIO(data)) as z,zipfile.ZipFile(out,'w') as dest:
            for item in z.infolist():
                content=z.read(item.filename)
                if item.filename.endswith('sheet1.xml'):content=content.replace(b'<c r="A2" t="inlineStr"><is><t>SW</t></is></c>',b'<c r="A2" t="s"><v>bad-index</v></c>')
                dest.writestr(item,content)
        self.assertEqual(self.client.post('/api/reporting/inventory-excel',content=out.getvalue()).status_code,400)

if __name__=='__main__':unittest.main()
