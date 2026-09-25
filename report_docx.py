"""Create an editable network design report from a planned project."""

from datetime import datetime, timezone
from io import BytesIO
import base64
import binascii

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


def _shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def _border(cell):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right"):
        element = OxmlElement(f"w:{edge}")
        element.set(qn("w:val"), "single")
        element.set(qn("w:color"), "D9D9D9")
        element.set(qn("w:sz"), "4")
        borders.append(element)


def _padding(cell):
    tc_pr = cell._tc.get_or_add_tcPr()
    margins = OxmlElement("w:tcMar")
    for edge in ("top", "start", "bottom", "end"):
        node = OxmlElement(f"w:{edge}")
        node.set(qn("w:w"), "95")
        node.set(qn("w:type"), "dxa")
        margins.append(node)
    tc_pr.append(margins)


def _table(doc, columns, widths, rows):
    table = doc.add_table(rows=1, cols=len(columns))
    table.autofit = False
    for index, (label, width) in enumerate(zip(columns, widths)):
        cell = table.rows[0].cells[index]
        cell.width = Inches(width)
        cell.text = label
    header = table.rows[0]
    header._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
    for cell in header.cells:
        _shade(cell, "E6EFF6")
        for run in cell.paragraphs[0].runs:
            run.bold = True
    for row_index, values in enumerate(rows):
        cells = table.add_row().cells
        for index, value in enumerate(values):
            cells[index].width = Inches(widths[index])
            cells[index].text = str(value)
            if row_index % 2:
                _shade(cells[index], "F7FAFC")
    for row in table.rows:
        row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
        for cell in row.cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            _border(cell)
            _padding(cell)
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                for run in paragraph.runs:
                    run.font.size = Pt(9)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def _technology_text(project, placement):
    descriptions = {
        "vPC": "The design plans Cisco NX-OS virtual port channel redundancy. Assign device pairs, peer links, keepalive paths and member port channels before deployment.",
        "MLAG": "The design plans Arista MLAG redundancy. Assign device pairs, peer links and downstream port channels before deployment.",
        "M-LAG": "The design plans Huawei M-LAG redundancy using DFS and Eth-Trunk. Assign device pairs, peer links and heartbeat paths before deployment.",
        "STP": "The design uses spanning tree for loop prevention. Engineer root placement, link roles and edge-port protection before deployment.",
        "EVPN/VXLAN": "The design plans an EVPN/VXLAN fabric. Engineer underlay reachability, VTEPs, VNI mappings and BGP EVPN adjacencies separately.",
        "SD-WAN": "The SD-WAN edge design does not yet include a selected router configuration. Complete the router details in Technology Workspaces and reopen the report.",
        "None": "A redundancy or overlay technology has not yet been selected.",
    }
    text = descriptions.get(project["technology"], "The selected technology requires detailed engineering before deployment.")
    if project["technology"] != "None":
        text += f" Planned placement: {placement} tier."
    return text


def build_report_docx(project):
    """Return a Word file in memory with planned topology and configuration."""
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.left_margin = section.right_margin = Inches(.7)
    section.top_margin = section.bottom_margin = Inches(.7)

    styles = doc.styles
    styles["Normal"].font.name = "Arial"
    styles["Normal"].font.size = Pt(10)
    styles["Normal"].paragraph_format.space_after = Pt(7)
    for style_name, size in (("Title", 21), ("Heading 1", 13)):
        style = styles[style_name]
        style.font.name = "Arial"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(14 if style_name == "Heading 1" else 0)
        style.paragraph_format.space_after = Pt(8)
    title_style = styles["Title"].element
    title_properties = title_style.pPr
    if title_properties is not None:
        border = title_properties.find(qn("w:pBdr"))
        if border is not None:
            title_properties.remove(border)
    styles["Subtitle"].font.name = "Arial"
    styles["Subtitle"].font.size = Pt(12)
    styles["Subtitle"].font.italic = False
    styles["Subtitle"].font.color.rgb = RGBColor(0, 0, 0)

    name = project["name"].strip()
    title = doc.add_paragraph(name, style="Title")
    title.paragraph_format.keep_with_next = True
    doc.add_paragraph("Network design and device inventory", style="Subtitle")
    roles = project.get("roles") or {}
    upper = roles.get("upper") or ("Core" if project["architecture"] == "core-access" else "Spine")
    lower = roles.get("lower") or ("Access" if project["architecture"] == "core-access" else "Leaf")
    vendor = "Huawei CloudEngine" if project["vendor"] == "Huawei_CE_SW" else project["vendor"]
    doc.add_paragraph(
        f"Draft report  |  {datetime.now(timezone.utc).strftime('%d %B %Y')}  |  {vendor}"
    )

    doc.add_heading("Project overview", level=1)
    architecture = ("Module topology" if project["architecture"] == "module" else
                    "Core Access" if project["architecture"] == "core-access" else "Spine Leaf")
    overview = (f"The planned {architecture} has {project['upperCount']} {upper.lower()} "
                + ("device" if project["upperCount"] == 1 else "devices")
                + (f" and {project['lowerCount']} {lower.lower()} "
                   + ("device" if project["lowerCount"] == 1 else "devices") if project["lowerCount"] else "")
                + f" and uses {project['technology']}.")
    doc.add_paragraph(overview)
    doc.add_paragraph(project["scope"].strip() or "Project scope and design intent await engineer input.")

    devices = {device["id"]: device for device in project["devices"]}
    active = [link for link in project["links"] if link["enabled"]]
    device_name = lambda device: device["hostname"] or f"{upper if device['tier'] == 'upper' else lower}-{device['index']:02d}"

    doc.add_heading("Planned topology", level=1)
    doc.add_paragraph(
        f"{upper} tier: " + ", ".join(device_name(d) for d in project["devices"] if d["tier"] == "upper")
    )
    if project["lowerCount"]:
        doc.add_paragraph(
            f"{lower} tier: " + ", ".join(device_name(d) for d in project["devices"] if d["tier"] == "lower")
        )
    topology_png = project.get("topologyPng", "")
    if topology_png:
        prefix = "data:image/png;base64,"
        if not topology_png.startswith(prefix):
            raise ValueError("Invalid topology image format")
        try:
            picture = base64.b64decode(topology_png[len(prefix):], validate=True)
        except binascii.Error as exc:
            raise ValueError("Invalid topology image encoding") from exc
        if len(picture) > 2250000 or not picture.startswith(b"\x89PNG\r\n\x1a\n"):
            raise ValueError("Invalid topology PNG")
        doc.add_picture(BytesIO(picture), width=Inches(6.9))
    doc.add_heading("Connection schedule", level=1)
    _table(
        doc,
        [f"{upper} device", "Port", f"{lower} device", "Port", "Connection / speed"],
        [1.55, 1.1, 1.55, 1.1, 1.65],
        [[device_name(devices[l["a"]]), l["upperPort"] or "TBD", device_name(devices[l["b"]]), l["lowerPort"] or "TBD", " · ".join(filter(None, [l.get("detail"), l["speed"]])) or "TBD"] for l in active]
        or [["No physical links modeled", "", "", "", ""]],
    )
    if project.get("specialLinks"):
        doc.add_heading("Special connections and control paths", level=2)
        _table(
            doc,
            ["Type", "Endpoint A / port", "Endpoint B / port", "Details"],
            [1.25, 1.8, 1.8, 2.1],
            [[link["kind"] + (" (logical)" if link["logical"] else ""),
              device_name(devices[link["a"]]) + " · " + (link["aPort"] or "TBD"),
              device_name(devices[link["b"]]) + " · " + (link["bPort"] or "TBD"), link["detail"]]
             for link in project["specialLinks"]],
        )
        doc.add_paragraph("Logical control paths identify endpoints; they do not imply a direct physical cable.")

    doc.add_heading("Technology approach", level=1)
    placement = upper if project["techPlacement"] == "upper" else lower
    doc.add_paragraph(
        f"The planned {project['technology']} design uses {vendor}. Device configuration proposals are included in the appendix."
        if project.get("configurations") else _technology_text(project, placement)
    )
    doc.add_paragraph("Design intent only; operational state has not been verified.")
    decisions = [item for item in project.get("parameters", []) if item["value"] != "Design boundary"]
    design_notes = [item for item in project.get("parameters", []) if item["value"] == "Design boundary"]
    if decisions:
        doc.add_heading("Technology decisions and network impact", level=2)
        doc.add_paragraph("Architecture and operational behavior reflect the selected design. Actual forwarding depends on peer configuration and device state.")
        for item in decisions:
            paragraph = doc.add_paragraph()
            paragraph.paragraph_format.keep_together = True
            paragraph.add_run(f"{item['label']} · {item['value']}\n").bold = True
            paragraph.add_run(item.get("impact") or "This saved project predates design explanations; reopen the Technology Workspace to regenerate this decision.")
    if design_notes:
        doc.add_heading("Design assumptions and implementation notes", level=2)
        for item in design_notes:
            paragraph = doc.add_paragraph()
            paragraph.paragraph_format.keep_together = True
            paragraph.add_run(f"{item['label']}\n").bold = True
            paragraph.add_run(item.get("impact") or "")

    doc.add_page_break()
    doc.add_heading("Device inventory", level=1)
    source = lambda value: "Verified from device" if value == "device" else "Engineer entry" if value == "manual" else "Awaiting assignment"
    _table(
        doc,
        ["Device", "Role", "Model", "Serial", "Source"],
        [1.6, .9, 1.5, 1.4, 1.7],
        [[device_name(d), upper if d["tier"] == "upper" else lower, d["model"] or ("External" if d.get("external") else "Awaiting assignment"), d["serial"] or ("External" if d.get("external") else "Awaiting assignment"),
          "External peer" if d.get("external") else f"{source(d['modelSource'])} / {source(d['serialSource'])}" if d["modelSource"] or d["serialSource"] else "Planned"]
         for d in project["devices"]],
    )

    doc.add_heading("Items to confirm", level=1)
    incomplete = sum(not d.get("external") and (not d["model"] or not d["serial"]) for d in project["devices"])
    missing_ports = sum(not l["upperPort"] or not l["lowerPort"] for l in active)
    doc.add_paragraph(
        (f"{incomplete} {'device still needs' if incomplete == 1 else 'devices still need'} a model or serial number. " if incomplete
         else "All device models and serial numbers are populated. ")
        + (f"{missing_ports} active {'link still needs' if missing_ports == 1 else 'links still need'} a port at one or both ends. " if missing_ports
           else "All active link endpoint ports are assigned. ")
        + "Compare device-sourced inventory with the intended bill of materials before closeout."
    )
    if project.get("configurations"):
        doc.add_page_break()
        doc.add_heading("Appendix · Planned device configurations", level=1)
        doc.add_paragraph(
            "Generated configurations reflect design inputs and have not been validated on a device. "
            "Review against the target model and software release before use."
        )
        configurations = {config["deviceId"]: config for config in project["configurations"]}
        configured_devices = [device for device in project["devices"] if device["id"] in configurations]
        for index, device in enumerate(configured_devices):
            config = configurations[device["id"]]
            doc.add_heading(device_name(device), level=2)
            doc.add_paragraph(config["source"])
            for line in config["text"].splitlines():
                paragraph = doc.add_paragraph(style="Normal")
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.line_spacing = 1.0
                run = paragraph.add_run(line or " ")
                run.font.name = "Consolas"
                run.font.size = Pt(8)
            if index < len(configured_devices) - 1:
                doc.add_page_break()
    doc.core_properties.title = name
    doc.core_properties.author = "Network Configurator"
    output = BytesIO()
    doc.save(output)
    output.seek(0)
    return output
